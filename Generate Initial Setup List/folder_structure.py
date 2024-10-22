import os
from docx import Document

def list_folder_structure(folder_path, parent_doc):
    """
    Recursively traverses the folder structure and writes to the Word document.
    """
    for root, dirs, files in os.walk(folder_path):
        # Calculate the indentation level based on the folder depth
        level = root.replace(folder_path, "").count(os.sep)
        indent = "  " * level  # Two spaces per level

        # Add the directory name to the document
        parent_doc.add_paragraph(f"{indent}📂 {os.path.basename(root)}", style='Heading2')

        # List all files in the current directory
        for file in files:
            parent_doc.add_paragraph(f"{indent}  📄 {file}", style='BodyText')

def main():
    # Get the folder path from the user
    folder_path = input("Enter the folder path: ").strip()

    # Verify the path exists
    if not os.path.exists(folder_path):
        print(f"Error: The folder path '{folder_path}' does not exist.")
        return

    # Create a new Word document
    doc = Document()
    doc.add_heading("Folder Structure Report", 0)

    # List folder structure in the Word document
    list_folder_structure(folder_path, doc)

    # Save the document in the same directory as the Python script
    output_path = os.path.join(os.path.dirname(__file__), "Folder_Structure_Report.docx")
    doc.save(output_path)

    print(f"Report generated successfully: {output_path}")

if __name__ == "__main__":
    main()
